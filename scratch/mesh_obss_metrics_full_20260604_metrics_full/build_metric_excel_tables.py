#!/usr/bin/env python3
"""Generate readable Excel tables from mesh OBSS metrics CSV files."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


RESULT_RE = re.compile(r"^mode(?P<mode>\d+)(?:_A(?P<link_a>[^_]+)_B(?P<link_b>[^_]+))?_(?P<assoc>ont|ap1|ap2)\.csv$")
ASSOC_ORDER = {"ont": 0, "ap1": 1, "ap2": 2}
INT_RE = re.compile(r"^[+-]?\d+$")
FLOAT_RE = re.compile(r"^[+-]?(?:\d+\.\d*|\d*\.\d+)(?:[eE][+-]?\d+)?$|^[+-]?\d+[eE][+-]?\d+$")


SCOPE_ROWS = [
    ("基础维度", "模式", "mode", "1..8", "拓扑模式编号。", "1-4 为星型 L1，5-8 为链型 L2。"),
    ("基础维度", "场景", "scenario", "字符串", "模式名称，形如 L1_STAR_WW、L2_CHAIN_EE。", "W/E 分别表示无线/有线回传。"),
    ("基础维度", "拓扑", "topology", "star/chain", "由场景名解析得到。", "-"),
    ("基础维度", "链路A类型", "linkAType", "wireless/wired", "由场景名末尾第 1 个 W/E 解析得到。", "-"),
    ("基础维度", "链路B类型", "linkBType", "wireless/wired", "由场景名末尾第 2 个 W/E 解析得到。", "-"),
    ("基础维度", "STA接入", "staAssoc", "ont/ap1/ap2", "STA 当前接入的基础设施节点。", "决定下行路径和最后一跳。"),
    ("基础维度", "X坐标/Y坐标", "x,y", "m", "STA 网格点坐标。", "-"),
    ("基础维度", "平均次数", "runCount", "count", "同一点位参与平均的 seed 数。", "-"),
    ("基础维度", "业务类型", "trafficType", "tcp/udp", "主下行业务流量类型。", "本轮用于区分 TCP 与 UDP。"),
    ("总吞吐", "E2E吞吐", "endToEndMbps", "Mbps", "测量窗口内应用层 sink 接收字节换算的端到端吞吐。", "TCP/UDP 通用接收口径。"),
    ("分跳吞吐", "HN链路", "hopNName", "字符串", "第 N 条活动转发链路名称。", "最多 3 跳。"),
    ("分跳吞吐", "HN吞吐", "hopNMbps", "Mbps", "第 N 跳接收侧 IPv4 Rx 前向 TCP 数据包吞吐。", "通常略高于应用层吞吐。"),
    ("分跳时延", "HN RTT avg", "hopNRttAvgMs", "ms", "第 N 跳 UDP echo probe RTT 平均值。", "主动探针，不是 TCP 端到端 RTT。"),
    ("分跳时延", "HN RTT p95", "hopNRttP95Ms", "ms", "第 N 跳 UDP echo probe RTT 95 分位值。", "尾延迟口径。"),
    ("分跳时延", "HN RTT loss", "hopNRttLoss", "比例", "第 N 跳 UDP echo probe 未收到 echo 的比例。", "近似 loss/timeout。"),
    ("分跳时延", "HN RTT jitter", "hopNRttJitterMs", "ms", "第 N 跳 UDP echo probe RTT 样本标准差。", "反映时延波动。"),
    ("链路类型", "HN无线标志", "hopNWireless", "0/1", "该跳是否为 Wi-Fi 链路。", "0 为有线 CSMA。"),
    ("速率/MCS", "HN MCS", "hopNMcsAvg", "MCS index", "匹配下行数据 PSDU 的 MCS 平均值。", "-"),
    ("速率/MCS", "HN PHY速率", "hopNPhyRateMbpsAvg", "Mbps", "匹配下行数据 PSDU 的 PHY rate 平均值。", "-"),
    ("速率/MCS", "HN数据速率", "hopNDataRateMbpsAvg", "Mbps", "匹配下行数据 PSDU 的 data rate 平均值。", "-"),
    ("空口占用", "HN空口占用", "hopNTxAirtimeDuty", "比例", "该跳下行数据 PSDU 发送时长累计 / test。", "不含 ACK/其它链路帧。"),
    ("聚合", "HN PPDU数", "hopNTxPpdus", "count", "测量窗口内匹配的下行 PPDU 数。", "-"),
    ("聚合", "HN MPDU数", "hopNTxMpdus", "count", "测量窗口内匹配的下行 MPDU 数。", "-"),
    ("聚合", "HN MPDU/PPDU", "hopNAvgMpdusPerPpdu", "count", "TxMpdus / TxPpdus。", "A-MPDU 聚合程度。"),
    ("聚合", "HN PSDU字节", "hopNAvgPsduBytes", "bytes", "TX PSDU 字节数累计 / TxPpdus。", "-"),
    ("聚合", "HN MSDU/MPDU", "hopNAvgMsdusPerMpdu", "count", "A-MSDU 中 MSDU 数累计 / TxMpdus。", "A-MSDU 聚合程度。"),
    ("接收质量", "HN MPDU PSR", "hopNRxMpduPsr", "比例", "成功 MPDU / 总 MPDU。", "-"),
    ("接收质量", "HN MPDU PER", "hopNRxMpduPer", "比例", "失败 MPDU / 总 MPDU。", "-"),
    ("接收质量", "HN PPDU PSR", "hopNRxPpduPsr", "比例", "未出现 MPDU 失败的 PPDU / 接收 PPDU。", "-"),
    ("接收质量", "HN RSSI", "hopNRxRssiDbmAvg", "dBm", "匹配下行数据帧的 RSSI 平均值。", "-"),
    ("接收质量", "HN SNR", "hopNRxSnrDbAvg", "dB", "匹配下行数据帧的 SNR 平均值。", "-"),
    ("TCP", "TCP ACK PPS", "tcpAckPktsPerSec", "pkt/s", "ONT TCP socket ACK 包数 / test。", "-"),
    ("TCP", "TCP ACK Bps", "tcpAckBytesPerSec", "B/s", "ACK payload + TCP header 字节 / test。", "-"),
    ("TCP", "TCP数据PPS", "tcpDataPktsPerSec", "pkt/s", "ONT TCP socket payload>0 数据包数 / test。", "-"),
    ("TCP", "TCP数据吞吐", "tcpDataMbps", "Mbps", "源端 TCP dataBytes 换算吞吐。", "含 TCP header。"),
    ("TCP", "TCP重传数", "tcpRetransPkts", "count", "ONT TCP socket Retransmission trace 计数。", "-"),
    ("TCP", "TCP RTT", "tcpRttMsAvg", "ms", "TCP 端到端 RTT trace 平均值。", "不是每跳 RTT。"),
    ("TCP", "TCP LastRTT", "tcpLastRttMsAvg", "ms", "TCP LastRTT trace 平均值。", "-"),
    ("OBSS", "OBSS帧数", "nodeObssFrames", "count", "节点监听到 OBSS AP 发出的帧数。", "节点包括 ONT/AP1/AP2/STA。"),
    ("OBSS", "OBSS RSSI", "nodeObssRssiDbmAvg", "dBm", "节点监听到 OBSS 帧的 RSSI 平均值。", "无帧时为空，Excel 显示为 -。"),
    ("OBSS", "OBSS SNR", "nodeObssSnrDbAvg", "dB", "节点监听到 OBSS 帧的 SNR 平均值。", "无帧时为空，Excel 显示为 -。"),
    ("OBSS", "OBSS空口占用", "nodeObssAirtimeDuty", "比例", "节点看到的 OBSS PPDU 持续时间累计 / test。", "观察节点口径。"),
    ("OBSS", "OBSS PPDU PSR", "nodeObssPpduPsr", "比例", "节点看到的 OBSS PPDU 成功率。", "未出现 MPDU 失败的 PPDU 视为成功。"),
    ("OBSS", "OBSS实测占用", "obssMeasuredDuty", "比例", "OBSS AP PHY TxBegin/TxEnd 发送时长累计 / test。", "全局干扰源发送 duty。"),
    ("OBSS", "OBSS目标占用", "obssTargetDuty", "比例", "配置的目标 OBSS 空口占用。", "-"),
    ("OBSS", "OBSS速率", "obssRate", "DataRate", "OBSS UDP offered rate 配置。", "-"),
]


FIELD_ORDER = ["mode", "scenario", "topology", "linkAType", "linkBType", "staAssoc", "x", "y", "runCount", "trafficType", "endToEndMbps"]

for hop in (1, 2, 3):
    FIELD_ORDER.extend(
        [
            f"hop{hop}Name",
            f"hop{hop}Wireless",
            f"hop{hop}Mbps",
            f"hop{hop}RttAvgMs",
            f"hop{hop}RttP95Ms",
            f"hop{hop}RttLoss",
            f"hop{hop}RttJitterMs",
            f"hop{hop}McsAvg",
            f"hop{hop}PhyRateMbpsAvg",
            f"hop{hop}DataRateMbpsAvg",
            f"hop{hop}TxAirtimeDuty",
            f"hop{hop}TxPpdus",
            f"hop{hop}TxMpdus",
            f"hop{hop}AvgMpdusPerPpdu",
            f"hop{hop}AvgPsduBytes",
            f"hop{hop}AvgMsdusPerMpdu",
            f"hop{hop}RxMpduPsr",
            f"hop{hop}RxMpduPer",
            f"hop{hop}RxPpduPsr",
            f"hop{hop}RxRssiDbmAvg",
            f"hop{hop}RxSnrDbAvg",
        ]
    )

FIELD_ORDER.extend(
    [
        "tcpAckPktsPerSec",
        "tcpAckBytesPerSec",
        "tcpDataPktsPerSec",
        "tcpDataMbps",
        "tcpRetransPkts",
        "tcpRttMsAvg",
        "tcpLastRttMsAvg",
    ]
)

for node in ("ont", "ap1", "ap2", "sta"):
    FIELD_ORDER.extend(
        [
            f"{node}ObssFrames",
            f"{node}ObssRssiDbmAvg",
            f"{node}ObssSnrDbAvg",
            f"{node}ObssAirtimeDuty",
            f"{node}ObssPpduPsr",
        ]
    )

FIELD_ORDER.extend(["obssMeasuredDuty", "obssTargetDuty", "obssRate"])


PARAM_NAMES = {
    "mode": "模式",
    "scenario": "场景",
    "topology": "拓扑",
    "linkAType": "链路A类型",
    "linkBType": "链路B类型",
    "staAssoc": "STA接入",
    "x": "X坐标",
    "y": "Y坐标",
    "runCount": "平均次数",
    "trafficType": "业务类型",
    "endToEndMbps": "E2E吞吐(Mbps)",
    "tcpAckPktsPerSec": "TCP ACK PPS",
    "tcpAckBytesPerSec": "TCP ACK Bps",
    "tcpDataPktsPerSec": "TCP数据PPS",
    "tcpDataMbps": "TCP数据吞吐(Mbps)",
    "tcpRetransPkts": "TCP重传数",
    "tcpRttMsAvg": "TCP RTT(ms)",
    "tcpLastRttMsAvg": "TCP LastRTT(ms)",
    "obssMeasuredDuty": "OBSS实测占用",
    "obssTargetDuty": "OBSS目标占用",
    "obssRate": "OBSS速率",
}

for hop in (1, 2, 3):
    prefix = f"H{hop}"
    PARAM_NAMES.update(
        {
            f"hop{hop}Name": f"{prefix}链路",
            f"hop{hop}Wireless": f"{prefix}无线标志",
            f"hop{hop}Mbps": f"{prefix}吞吐(Mbps)",
            f"hop{hop}RttAvgMs": f"{prefix} RTT avg(ms)",
            f"hop{hop}RttP95Ms": f"{prefix} RTT p95(ms)",
            f"hop{hop}RttLoss": f"{prefix} RTT loss",
            f"hop{hop}RttJitterMs": f"{prefix} RTT jitter(ms)",
            f"hop{hop}McsAvg": f"{prefix} MCS",
            f"hop{hop}PhyRateMbpsAvg": f"{prefix} PHY速率(Mbps)",
            f"hop{hop}DataRateMbpsAvg": f"{prefix}数据速率(Mbps)",
            f"hop{hop}TxAirtimeDuty": f"{prefix}空口占用",
            f"hop{hop}TxPpdus": f"{prefix} PPDU数",
            f"hop{hop}TxMpdus": f"{prefix} MPDU数",
            f"hop{hop}AvgMpdusPerPpdu": f"{prefix} MPDU/PPDU",
            f"hop{hop}AvgPsduBytes": f"{prefix} PSDU字节",
            f"hop{hop}AvgMsdusPerMpdu": f"{prefix} MSDU/MPDU",
            f"hop{hop}RxMpduPsr": f"{prefix} MPDU PSR",
            f"hop{hop}RxMpduPer": f"{prefix} MPDU PER",
            f"hop{hop}RxPpduPsr": f"{prefix} PPDU PSR",
            f"hop{hop}RxRssiDbmAvg": f"{prefix} RSSI(dBm)",
            f"hop{hop}RxSnrDbAvg": f"{prefix} SNR(dB)",
        }
    )

for node, label in (("ont", "ONT"), ("ap1", "AP1"), ("ap2", "AP2"), ("sta", "STA")):
    PARAM_NAMES.update(
        {
            f"{node}ObssFrames": f"{label} OBSS帧数",
            f"{node}ObssRssiDbmAvg": f"{label} OBSS RSSI(dBm)",
            f"{node}ObssSnrDbAvg": f"{label} OBSS SNR(dB)",
            f"{node}ObssAirtimeDuty": f"{label} OBSS空口占用",
            f"{node}ObssPpduPsr": f"{label} OBSS PPDU PSR",
        }
    )


def clean_cell(value: object) -> str:
    if value is None:
        return "-"
    text = str(value).strip()
    return text if text else "-"


def excel_value(value: object) -> object:
    text = clean_cell(value)
    if text == "-":
        return text
    if INT_RE.match(text):
        return int(text)
    if FLOAT_RE.match(text):
        return float(text)
    return text


def scenario_parts(scenario: str) -> tuple[str, str, str]:
    parts = scenario.split("_")
    topology = "chain" if "CHAIN" in parts else "star" if "STAR" in parts else "-"
    suffix = parts[-1] if parts else ""
    link_a = "wireless" if len(suffix) >= 1 and suffix[0] == "W" else "wired" if len(suffix) >= 1 and suffix[0] == "E" else "-"
    link_b = "wireless" if len(suffix) >= 2 and suffix[1] == "W" else "wired" if len(suffix) >= 2 and suffix[1] == "E" else "-"
    return topology, link_a, link_b


def result_csvs(input_dir: Path) -> list[Path]:
    files = [path for path in input_dir.glob("mode*.csv") if RESULT_RE.match(path.name)]
    def key(path: Path):
        match = RESULT_RE.match(path.name)
        return (
            int(match.group("mode")),
            match.group("link_a") or "",
            match.group("link_b") or "",
            ASSOC_ORDER[match.group("assoc")],
        )
    return sorted(files, key=key)


def read_theory_rows(input_dir: Path) -> list[list[object]]:
    rows = [[PARAM_NAMES[field] for field in FIELD_ORDER]]
    for path in result_csvs(input_dir):
        with path.open(newline="", encoding="utf-8-sig") as stream:
            reader = csv.DictReader(stream)
            for source in reader:
                topology, link_a, link_b = scenario_parts(clean_cell(source.get("scenario")))
                generated = {"topology": topology, "linkAType": link_a, "linkBType": link_b}
                row = []
                for field in FIELD_ORDER:
                    value = source.get(field)
                    if clean_cell(value) == "-" and field in generated:
                        value = generated[field]
                    row.append(excel_value(value))
                rows.append(row)
    if len(rows) == 1:
        raise SystemExit(f"No modeN_ont/ap1/ap2.csv files found in {input_dir}")
    return rows


def scope_rows() -> list[list[object]]:
    return [["指标组", "参数名/简称", "字段", "单位/取值", "统计口径", "备注"]] + [list(row) for row in SCOPE_ROWS]


def add_sheet(wb: Workbook, title: str, rows: list[list[object]], table_name: str) -> None:
    ws = wb.create_sheet(title=title)
    for row in rows:
        ws.append([excel_value(value) for value in row])

    max_row = ws.max_row
    max_col = ws.max_column
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        sample_end = min(max_row, 100)
        width = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, sample_end + 1)) + 2
        ws.column_dimensions[letter].width = min(max(width, 8), 36)

    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(max_col)}{max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def write_workbook(path: Path, sheets: Iterable[tuple[str, list[list[object]], str]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows, table_name in sheets:
        add_sheet(wb, title, rows, table_name)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Excel metric tables from mesh OBSS CSV results.")
    parser.add_argument("--input-dir", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--output-dir", type=Path, default=None)
    args = parser.parse_args()

    input_dir = args.input_dir.resolve()
    output_dir = (args.output_dir or args.input_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    scope = scope_rows()
    theory = read_theory_rows(input_dir)
    write_workbook(output_dir / "metric_statistical_scope.xlsx", [("指标统计口径", scope, "MetricScopeTable")])
    write_workbook(output_dir / "theory_reference_per_point_metrics.xlsx", [("理论参考宽表", theory, "TheoryReferenceTable")])
    write_workbook(
        output_dir / "mesh_obss_metrics_full_tables.xlsx",
        [
            ("指标统计口径", scope, "MetricScopeCombinedTable"),
            ("理论参考宽表", theory, "TheoryReferenceCombinedTable"),
        ],
    )

    print(f"input_dir={input_dir}")
    print(f"output_dir={output_dir}")
    print(f"metric_scope_rows={len(scope) - 1}, metric_scope_cols={len(scope[0])}")
    print(f"theory_rows={len(theory) - 1}, theory_cols={len(theory[0])}")


if __name__ == "__main__":
    main()
